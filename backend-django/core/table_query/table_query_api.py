#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Table Query API - 表查询管理接口
提供动态表查询、配置管理和数据导出功能

安全措施：
- 表名白名单验证
- 字段白名单验证
- 参数化查询防止 SQL 注入
- 排序字段验证
- 操作符白名单
"""
import io
import logging
import re
import time
from typing import List, Optional, Tuple, Any

from django.db import connection
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from ninja import Router, Query
from ninja.errors import HttpError
from ninja.pagination import paginate

from common.fu_crud import create, retrieve, delete, update
from common.fu_pagination import MyPagination
from core.table_query.table_query_model import TableQueryConfig, TableQueryLog
from core.table_query.table_query_schema import (
    TableQueryConfigSchemaIn,
    TableQueryConfigSchemaPatch,
    TableQueryConfigSchemaOut,
    TableQueryConfigFilters,
    TableQueryIn,
    TableQueryResult,
    ExportParams,
    FilterCondition,
    TableQueryLogSchemaOut,
    TableQueryLogFilters,
    JoinQueryIn,
    JoinQueryResult,
    JoinPreviewOut,
    JoinExportParams,
    JoinInfo,
    JoinTableInfo,
    FieldInfo,
)

logger = logging.getLogger(__name__)
router = Router()

# =============================================================================
# 安全验证相关常量
# =============================================================================

# 允许的 SQL 操作符
ALLOWED_OPERATORS = {
    "eq": "=",           # 等于
    "ne": "!=",          # 不等于
    "gt": ">",           # 大于
    "gte": ">=",         # 大于等于
    "lt": "<",           # 小于
    "lte": "<=",         # 小于等于
    "like": "LIKE",      # 模糊匹配
    "in": "IN",          # 包含
    "between": "BETWEEN", # 范围
}

# 表名和字段名的合法字符正则
IDENTIFIER_PATTERN = re.compile(r'^[a-zA-Z_][a-zA-Z0-9_]*$')


def quote_identifier(name: str) -> str:
    """
    使用反引号转义标识符（表名/字段名）
    防止 MySQL 保留字冲突（如 partition, explain, order, key 等）
    
    Args:
        name: 标识符名称
    
    Returns:
        反引号包围的标识符
    """
    # 移除可能存在的反引号，避免重复转义
    name = name.replace('`', '')
    return f'`{name}`'


# =============================================================================
# 安全验证函数
# =============================================================================

def validate_identifier(name: str, identifier_type: str = "identifier") -> str:
    """
    验证标识符（表名/字段名）是否合法
    
    Args:
        name: 标识符名称
        identifier_type: 标识符类型（用于错误信息）
    
    Returns:
        验证通过的标识符
    
    Raises:
        HttpError: 标识符包含非法字符
    """
    if not name or not IDENTIFIER_PATTERN.match(name):
        raise HttpError(400, f"非法的{identifier_type}: {name}")
    return name


def validate_table_name(table_name: str, config: TableQueryConfig) -> str:
    """
    验证表名是否在白名单中
    
    Args:
        table_name: 要验证的表名
        config: 配置对象
    
    Returns:
        验证通过的表名
    
    Raises:
        HttpError: 表名不在白名单中
    """
    # 首先验证格式
    validate_identifier(table_name, "表名")
    
    # 检查是否与配置中的表名匹配
    if table_name.upper() != config.table_name.upper():
        raise HttpError(403, f"表名不匹配: {table_name}")
    
    return config.table_name


def validate_fields(fields: Optional[List[str]], config: TableQueryConfig) -> List[str]:
    """
    验证字段列表是否在白名单中
    
    Args:
        fields: 要验证的字段列表
        config: 配置对象
    
    Returns:
        验证通过的字段列表
    
    Raises:
        HttpError: 字段不在白名单中
    """
    # 获取配置中的所有字段
    config_fields = config.config_json.get('fields', [])
    allowed_fields = {f['name'].upper() for f in config_fields}
    
    if not fields:
        # 返回所有可见字段
        return [f['name'] for f in config_fields if f.get('visible', True)]
    
    # 验证请求的字段
    validated_fields = []
    for field in fields:
        validate_identifier(field, "字段名")
        if field.upper() not in allowed_fields:
            raise HttpError(400, f"字段 {field} 不在允许列表中")
        validated_fields.append(field)
    
    return validated_fields


def validate_order_by(order_by: Optional[str], config: TableQueryConfig) -> str:
    """
    验证排序字段是否在白名单中
    
    Args:
        order_by: 排序字符串，如 "name ASC, id DESC"
        config: 配置对象
    
    Returns:
        验证通过的排序字符串
    
    Raises:
        HttpError: 排序字段不在白名单中
    """
    if not order_by:
        # 支持 camelCase 和 snake_case 两种格式
        default = config.config_json.get('defaultOrderBy') or config.config_json.get('default_order_by')
        if not default:
            # 使用第一个可排序字段作为默认排序
            fields = config.config_json.get('fields', [])
            sortable = [f['name'] for f in fields if f.get('sortable', False)]
            default = f"{sortable[0]} DESC" if sortable else None
        return default
    
    # 获取可排序字段
    sortable_fields = {f['name'].upper() for f in config.config_json.get('fields', []) 
                       if f.get('sortable', False)}
    
    # 解析排序字符串
    parts = []
    for part in order_by.split(','):
        part = part.strip()
        if not part:
            continue
        
        # 提取字段名和排序方向
        tokens = part.split()
        if not tokens:
            continue
        
        field_name = tokens[0]
        direction = tokens[1].upper() if len(tokens) > 1 else 'ASC'
        
        # 验证字段名
        validate_identifier(field_name, "排序字段")
        if field_name.upper() not in sortable_fields:
            raise HttpError(400, f"字段 {field_name} 不支持排序")
        
        # 验证排序方向
        if direction not in ('ASC', 'DESC'):
            raise HttpError(400, f"非法的排序方向: {direction}")
        
        parts.append(f"{quote_identifier(field_name)} {direction}")
    
    if parts:
        return ', '.join(parts)
    # 回退到默认排序
    default = config.config_json.get('defaultOrderBy') or config.config_json.get('default_order_by')
    if default:
        # 对默认排序字段也进行转义
        default_parts = []
        for part in default.split(','):
            part = part.strip()
            tokens = part.split()
            if tokens:
                field_name = tokens[0]
                direction = tokens[1].upper() if len(tokens) > 1 else 'DESC'
                default_parts.append(f"{quote_identifier(field_name)} {direction}")
        return ', '.join(default_parts) if default_parts else None
    else:
        fields = config.config_json.get('fields', [])
        sortable = [f['name'] for f in fields if f.get('sortable', False)]
        return f"{quote_identifier(sortable[0])} DESC" if sortable else None


def build_where_clause(
    filters: Optional[List[FilterCondition]], 
    config: TableQueryConfig
) -> Tuple[List[str], List[Any]]:
    """
    安全构建 WHERE 子句
    
    Args:
        filters: 过滤条件列表
        config: 配置对象
    
    Returns:
        (where子句列表, 参数列表)
    
    Raises:
        HttpError: 字段或操作符不合法
    """
    if not filters:
        return [], []
    
    # 获取可搜索字段
    searchable_fields = {f['name'].upper(): f for f in config.config_json.get('fields', []) 
                         if f.get('searchable', False)}
    
    where_clauses = []
    params = []
    
    for condition in filters:
        field = condition.field
        operator = condition.operator
        value = condition.value
        
        # 验证字段名
        validate_identifier(field, "过滤字段")
        field_upper = field.upper()
        
        if field_upper not in searchable_fields:
            # 获取所有可搜索字段列表，供调用方参考
            available_fields = [
                {
                    "name": f['name'],
                    "display_name": f.get('displayName', f.get('display_name', f['name'])),
                    "type": f.get('type', 'string')
                }
                for f in config.config_json.get('fields', [])
                if f.get('searchable', False)
            ]
            # HttpError 需要字符串参数，将字典序列化为 JSON
            import json
            error_detail = json.dumps({
                "message": f"字段 {field} 不支持搜索",
                "available_fields": available_fields
            }, ensure_ascii=False)
            raise HttpError(400, error_detail)
        
        # 验证操作符
        if operator not in ALLOWED_OPERATORS:
            supported_ops = ", ".join([f"{k}({v})" for k, v in ALLOWED_OPERATORS.items()])
            raise HttpError(400, f"不支持的操作符: {operator}。支持的操作符: {supported_ops}")
        
        sql_operator = ALLOWED_OPERATORS[operator]
        
        # 根据操作符构建 WHERE 子句（使用反引号转义字段名）
        quoted_field = quote_identifier(field)
        if operator == "like":
            where_clauses.append(f"{quoted_field} {sql_operator} %s")
            params.append(f"%{value}%")
        elif operator == "in":
            if not isinstance(value, (list, tuple)):
                value = [value]
            placeholders = ', '.join(['%s'] * len(value))
            where_clauses.append(f"{quoted_field} IN ({placeholders})")
            params.extend(value)
        elif operator == "between":
            if not isinstance(value, (list, tuple)) or len(value) != 2:
                raise HttpError(400, "BETWEEN 操作需要两个值")
            where_clauses.append(f"{quoted_field} BETWEEN %s AND %s")
            params.extend(value)
        else:
            where_clauses.append(f"{quoted_field} {sql_operator} %s")
            params.append(value)
    
    return where_clauses, params


def get_total_count(table_name: str, where_clauses: List[str], params: List[Any]) -> int:
    """
    获取查询结果总数
    
    Args:
        table_name: 表名
        where_clauses: WHERE 子句列表
        params: 参数列表
    
    Returns:
        结果总数
    """
    sql = f"SELECT COUNT(*) FROM {quote_identifier(table_name)}"
    if where_clauses:
        sql += f" WHERE {' AND '.join(where_clauses)}"
    
    with connection.cursor() as cursor:
        cursor.execute(sql, params)
        return cursor.fetchone()[0]


def _resolve_config_by_name(config_name: str) -> TableQueryConfig:
    """
    根据配置名称解析配置（支持模糊匹配）
    
    查询逻辑:
    1. 优先返回显示名称完全匹配的配置
    2. 如果没有完全匹配，返回显示名称包含关键字的第一个配置
    3. 如果都没有匹配，抛出 HttpError
    
    Args:
        config_name: 配置名称
    
    Returns:
        TableQueryConfig: 匹配的配置
    
    Raises:
        HttpError: 未找到匹配的配置
    """
    # 优先精确匹配
    config = TableQueryConfig.objects.filter(
        display_name=config_name, 
        is_deleted=False
    ).first()
    if config:
        return config
    
    # 模糊匹配
    config = TableQueryConfig.objects.filter(
        display_name__icontains=config_name, 
        is_deleted=False
    ).first()
    if config:
        return config
    
    raise HttpError(400, f"未找到名称匹配 '{config_name}' 的表查询配置")


# =============================================================================
# 配置管理 API
# =============================================================================

@router.post("/table-query/configs", response=TableQueryConfigSchemaOut, tags=["表查询管理"], summary="创建表查询配置")
def create_config(request, data: TableQueryConfigSchemaIn):
    """
    创建表查询配置
    
    请求体:
    - table_name: 数据库表名 (必填，唯一)
    - display_name: 显示名称 (必填)
    - description: 表描述 (可选)
    - config_json: 配置内容 (可选)
    - is_active: 是否激活 (可选，默认 True)
    """
    instance = create(request, data, TableQueryConfig)
    logger.info(f"创建表查询配置: {instance.table_name}")
    return instance


@router.get("/table-query/configs", response=List[TableQueryConfigSchemaOut], tags=["表查询管理"], summary="获取表查询配置列表（分页）")
@paginate(MyPagination)
def list_configs(request, filters: TableQueryConfigFilters = Query(...)):
    """
    获取表查询配置列表 (分页)
    
    查询参数:
    - page: 页码 (可选，默认为 1)
    - page_size: 每页数量 (可选，默认为 20)
    - table_name: 表名 (可选，模糊查询)
    - display_name: 显示名称 (可选，模糊查询)
    - is_active: 是否激活 (可选)
    """
    return retrieve(request, TableQueryConfig, filters)


@router.get("/table-query/configs/all", response=List[TableQueryConfigSchemaOut], tags=["表查询管理"], summary="获取所有表查询配置")
def list_all_configs(request, is_active: Optional[bool] = True):
    """
    获取所有表查询配置 (不分页)
    
    查询参数:
    - is_active: 是否只返回激活的配置 (可选，默认 True)
    """
    queryset = TableQueryConfig.objects.filter(is_deleted=False)
    if is_active is not None:
        queryset = queryset.filter(is_active=is_active)
    return list(queryset)


@router.get("/table-query/configs/{config_id}", response=TableQueryConfigSchemaOut, tags=["表查询管理"], summary="获取表查询配置详情")
def get_config(request, config_id: str):
    """
    获取表查询配置详情
    
    路径参数:
    - config_id: 配置ID
    """
    return get_object_or_404(TableQueryConfig, id=config_id, is_deleted=False)


@router.get("/table-query/configs/by-name/{name}", response=TableQueryConfigSchemaOut, tags=["表查询管理"], summary="按名称查询配置")
def get_config_by_name(request, name: str):
    """
    按配置名称查询表查询配置（支持模糊匹配）
    
    路径参数:
    - name: 配置显示名称（支持模糊匹配）
    
    查询逻辑:
    1. 优先返回显示名称完全匹配的配置
    2. 如果没有完全匹配，返回显示名称包含关键字的第一个配置
    3. 如果都没有匹配，返回 404 错误
    
    AI 调用建议: 直接传入用户提到的配置名称，无需获取 ID。
    
    示例:
    - /table-query/configs/by-name/用户表查询 → 精确匹配
    - /table-query/configs/by-name/用户 → 模糊匹配到"用户表查询"
    """
    # 优先精确匹配
    config = TableQueryConfig.objects.filter(display_name=name, is_deleted=False).first()
    if config:
        return config
    
    # 模糊匹配
    config = TableQueryConfig.objects.filter(display_name__icontains=name, is_deleted=False).first()
    if config:
        return config
    
    raise HttpError(404, f"未找到名称匹配 '{name}' 的表查询配置")


@router.get("/table-query/configs/{config_id}/searchable-fields", tags=["表查询管理"], summary="获取配置的可搜索字段")
def get_searchable_fields(request, config_id: str):
    """
    获取指定配置的可搜索字段列表
    
    路径参数:
    - config_id: 配置ID
    
    返回:
    - config_name: 配置显示名称
    - table_name: 数据库表名
    - searchable_fields: 可搜索字段列表，每个字段包含 name、display_name、type
    
    AI 调用建议: 在生成带过滤条件的查询脚本前，先调用此接口获取可用的过滤字段。
    """
    config = get_object_or_404(TableQueryConfig, id=config_id, is_deleted=False)
    fields = config.config_json.get('fields', [])
    searchable = [
        {
            "name": f['name'],
            "display_name": f.get('displayName', f.get('display_name', f['name'])),
            "type": f.get('type', 'string')
        }
        for f in fields if f.get('searchable', False)
    ]
    return {
        "config_name": config.display_name,
        "table_name": config.table_name,
        "searchable_fields": searchable
    }


@router.get("/table-query/configs/by-name/{name}/searchable-fields", tags=["表查询管理"], summary="按名称获取可搜索字段")
def get_searchable_fields_by_name(request, name: str):
    """
    按配置名称获取可搜索字段列表（支持模糊匹配）
    
    路径参数:
    - name: 配置显示名称（支持模糊匹配）
    
    返回:
    - config_name: 配置显示名称
    - table_name: 数据库表名
    - searchable_fields: 可搜索字段列表
    
    AI 调用建议: 直接传入用户提到的配置名称，无需先获取 ID。
    """
    # 优先精确匹配
    config = TableQueryConfig.objects.filter(display_name=name, is_deleted=False).first()
    if not config:
        # 模糊匹配
        config = TableQueryConfig.objects.filter(display_name__icontains=name, is_deleted=False).first()
    
    if not config:
        raise HttpError(404, f"未找到名称匹配 '{name}' 的表查询配置")
    
    fields = config.config_json.get('fields', [])
    searchable = [
        {
            "name": f['name'],
            "display_name": f.get('displayName', f.get('display_name', f['name'])),
            "type": f.get('type', 'string')
        }
        for f in fields if f.get('searchable', False)
    ]
    return {
        "config_name": config.display_name,
        "table_name": config.table_name,
        "searchable_fields": searchable
    }


@router.put("/table-query/configs/{config_id}", response=TableQueryConfigSchemaOut, tags=["表查询管理"], summary="更新表查询配置")
def update_config(request, config_id: str, data: TableQueryConfigSchemaPatch):
    """
    更新表查询配置
    
    路径参数:
    - config_id: 配置ID
    """
    instance = update(request, config_id, data, TableQueryConfig)
    logger.info(f"更新表查询配置: {instance.table_name}")
    return instance


@router.delete("/table-query/configs/{config_id}", response=TableQueryConfigSchemaOut, tags=["表查询管理"], summary="删除表查询配置")
def delete_config(request, config_id: str):
    """
    删除表查询配置 (软删除)
    
    路径参数:
    - config_id: 配置ID
    """
    instance = delete(config_id, TableQueryConfig)
    logger.info(f"删除表查询配置: {instance.table_name}")
    return instance


# =============================================================================
# 动态查询 API
# =============================================================================

@router.post("/table-query/query", response=TableQueryResult, tags=["表查询管理"], summary="执行动态表查询")
def execute_query(request, data: TableQueryIn):
    """
    执行动态表查询，支持分页、过滤、排序
    
    支持两种方式指定配置：
    1. config_id: 配置 ID（精确匹配）
    2. config_name: 配置名称（支持模糊匹配）
    
    参数优先级：config_id > config_name
    
    AI 调用建议：优先使用 config_name 参数，传入用户提到的配置名称即可。
    
    请求体:
    - config_id: 配置ID (可选，优先级高于 config_name)
    - config_name: 配置名称 (可选，支持模糊匹配，AI 推荐使用)
    - page: 页码 (可选，默认 1)
    - page_size: 每页数量 (可选，默认 20)
    - fields: 要查询的字段 (可选，默认为所有可见字段)
    - filters: 过滤条件 (可选)
    - order_by: 排序 (可选)
    
    安全措施:
    - 表名白名单验证
    - 字段白名单验证
    - 参数化查询防止 SQL 注入
    """
    start_time = time.time()
    
    # 获取配置（优先使用 config_id，其次使用 config_name）
    if data.config_id:
        config = get_object_or_404(TableQueryConfig, id=data.config_id, is_deleted=False)
    elif data.config_name:
        config = _resolve_config_by_name(data.config_name)
    else:
        raise HttpError(400, "必须提供 config_id 或 config_name 参数")
    
    if not config.is_active:
        raise HttpError(400, "该表查询配置已禁用")
    
    # 检查是否允许查询操作
    allowed_ops = config.config_json.get('allowedOperations') or config.config_json.get('allowed_operations', ['query'])
    if 'query' not in allowed_ops:
        raise HttpError(403, "此配置不允许查询操作")
    
    # 验证并获取字段
    select_fields = validate_fields(data.fields, config)
    
    # 构建安全的 WHERE 子句
    where_clauses, params = build_where_clause(data.filters, config)
    
    # 验证排序字段
    order_by = validate_order_by(data.order_by, config)
    
    # 分页限制
    max_page_size = config.config_json.get('maxPageSize') or config.config_json.get('max_page_size', 100)
    page_size = min(data.page_size, max_page_size)
    offset = (data.page - 1) * page_size
    
    # 构建查询 SQL（使用反引号转义表名和字段名）
    table_name = config.table_name
    quoted_fields = [quote_identifier(f) for f in select_fields]
    sql = f"SELECT {', '.join(quoted_fields)} FROM {quote_identifier(table_name)}"
    if where_clauses:
        sql += f" WHERE {' AND '.join(where_clauses)}"
    sql += f" ORDER BY {order_by}"
    sql += " LIMIT %s OFFSET %s"
    
    query_params = params + [page_size, offset]
    
    # 执行查询
    with connection.cursor() as cursor:
        cursor.execute(sql, query_params)
        columns = [col[0] for col in cursor.description]
        results = [dict(zip(columns, row)) for row in cursor.fetchall()]
    
    # 获取总数
    total = get_total_count(table_name, where_clauses, params)
    
    # 计算执行时间
    execution_time = (time.time() - start_time) * 1000
    
    # 记录查询日志
    user_id = str(request.auth.id) if hasattr(request, 'auth') and request.auth else 'anonymous'
    TableQueryLog.objects.create(
        user_id=user_id,
        table_name=table_name,
        operation="query",
        filters={
            "conditions": [f.dict() for f in (data.filters or [])],
            "order_by": order_by,
        },
        record_count=len(results),
        execution_time=execution_time,
    )
    
    logger.info(f"表查询执行: {table_name}, 返回 {len(results)} 条, 耗时 {execution_time:.2f}ms")
    
    return TableQueryResult(
        items=results,
        total=total,
        page=data.page,
        page_size=page_size,
    )


# =============================================================================
# 数据导出 API
# =============================================================================

@router.post("/table-query/export", tags=["表查询管理"], summary="导出表数据")
def export_data(request, data: ExportParams):
    """
    导出表数据 (Excel/CSV)
    
    支持两种方式指定配置：
    1. config_id: 配置 ID（精确匹配）
    2. config_name: 配置名称（支持模糊匹配）
    
    参数优先级：config_id > config_name
    
    AI 调用建议：优先使用 config_name 参数，传入用户提到的配置名称即可。
    
    请求体:
    - config_id: 配置ID (可选，优先级高于 config_name)
    - config_name: 配置名称 (可选，支持模糊匹配，AI 推荐使用)
    - format: 导出格式 excel/csv (可选，默认 excel)
    - fields: 要导出的字段 (可选，默认为所有可见字段)
    - filters: 过滤条件 (可选)
    - max_rows: 最大导出行数 (可选，默认 10000)
    
    返回:
    - 文件下载响应
    """
    start_time = time.time()
    
    # 获取配置（优先使用 config_id，其次使用 config_name）
    if data.config_id:
        config = get_object_or_404(TableQueryConfig, id=data.config_id, is_deleted=False)
    elif data.config_name:
        config = _resolve_config_by_name(data.config_name)
    else:
        raise HttpError(400, "必须提供 config_id 或 config_name 参数")
    
    if not config.is_active:
        raise HttpError(400, "该表查询配置已禁用")
    
    # 检查是否允许导出操作
    allowed_ops = config.config_json.get('allowedOperations') or config.config_json.get('allowed_operations', ['query'])
    if 'export' not in allowed_ops:
        raise HttpError(403, "此配置不允许导出操作")
    
    # 验证并获取字段
    select_fields = validate_fields(data.fields, config)
    
    # 构建安全的 WHERE 子句
    where_clauses, params = build_where_clause(data.filters, config)
    
    # 构建查询 SQL（使用反引号转义表名和字段名）
    table_name = config.table_name
    quoted_fields = [quote_identifier(f) for f in select_fields]
    sql = f"SELECT {', '.join(quoted_fields)} FROM {quote_identifier(table_name)}"
    if where_clauses:
        sql += f" WHERE {' AND '.join(where_clauses)}"
    sql += f" LIMIT %s"
    
    query_params = params + [data.max_rows]
    
    # 执行查询
    with connection.cursor() as cursor:
        cursor.execute(sql, query_params)
        columns = [col[0] for col in cursor.description]
        rows = cursor.fetchall()
    
    # 获取字段显示名称
    field_display_names = {}
    for field_config in config.config_json.get('fields', []):
        field_display_names[field_config['name']] = field_config.get('display_name', field_config['name'])
    
    headers = [field_display_names.get(col, col) for col in columns]
    
    # 计算执行时间
    execution_time = (time.time() - start_time) * 1000
    
    # 记录导出日志
    user_id = str(request.auth.id) if hasattr(request, 'auth') and request.auth else 'anonymous'
    TableQueryLog.objects.create(
        user_id=user_id,
        table_name=table_name,
        operation="export",
        filters={
            "conditions": [f.dict() for f in (data.filters or [])],
            "format": data.format,
        },
        record_count=len(rows),
        execution_time=execution_time,
    )
    
    logger.info(f"数据导出: {table_name}, 导出 {len(rows)} 条, 格式 {data.format}")
    
    # 生成导出文件
    if data.format == "csv":
        return _export_csv(config.display_name, headers, rows)
    else:
        return _export_excel(config.display_name, headers, rows)


def _export_csv(filename: str, headers: List[str], rows: List[tuple]) -> HttpResponse:
    """导出为 CSV 格式"""
    import csv
    
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    
    response = HttpResponse(
        output.getvalue(),
        content_type='text/csv; charset=utf-8-sig'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    return response


def _export_excel(filename: str, headers: List[str], rows: List[tuple]) -> HttpResponse:
    """导出为 Excel 格式"""
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HttpError(500, "Excel 导出需要安装 openpyxl 库")
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = filename[:31]  # Excel 工作表名最长 31 字符
    
    # 写入表头
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = openpyxl.styles.Font(bold=True)
    
    # 写入数据
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # 调整列宽
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15
    
    # 保存到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response


# =============================================================================
# 查询日志 API
# =============================================================================

@router.get("/table-query/logs", response=List[TableQueryLogSchemaOut], tags=["表查询管理"], summary="获取查询日志列表（分页）")
@paginate(MyPagination)
def list_logs(request, filters: TableQueryLogFilters = Query(...)):
    """
    获取查询日志列表 (分页)
    
    查询参数:
    - page: 页码 (可选，默认为 1)
    - page_size: 每页数量 (可选，默认为 20)
    - user_id: 用户ID (可选，精确查询)
    - table_name: 表名 (可选，模糊查询)
    - operation: 操作类型 (可选)
    """
    return retrieve(request, TableQueryLog, filters)


# =============================================================================
# 联合查询 API
# =============================================================================

@router.post("/table-query/join-query", response=JoinQueryResult, tags=["联合查询"], summary="执行外键联合查询")
def execute_join_query(request, data: JoinQueryIn):
    """
    执行基于外键关系的多表联合查询
    
    根据主表的外键关系自动 LEFT JOIN 关联表，将数据合并成一个结果集输出。
    
    请求体:
    - primary_table: 主表名或配置名（必填，支持模糊匹配）
    - max_depth: 最大关联深度（可选，1-5级，默认2级）
    - include_tables: 指定要包含的关联表（可选，优先级高于 exclude_tables）
    - exclude_tables: 指定要排除的关联表（可选）
    - page: 页码（可选，默认1）
    - page_size: 每页数量（可选，默认20，最大100）
    - filters: 过滤条件（可选，字段名需使用 表名_字段名 格式）
    - order_by: 排序（可选，字段名需使用 表名_字段名 格式）
    
    返回:
    - items: 数据列表
    - total: 总数
    - page: 当前页码
    - page_size: 每页数量
    - join_info: 关联信息（主表、关联表列表、关联深度等）
    - field_info: 字段元信息列表
    
    AI 调用建议: 直接传入用户提到的表名，系统会自动解析外键关系并进行联合查询。
    """
    from core.table_query.join_query_utils import execute_join_query as do_join_query
    
    start_time = time.time()
    
    # 转换过滤条件格式
    filters = None
    if data.filters:
        filters = [{"field": f.field, "operator": f.operator, "value": f.value} for f in data.filters]
    
    try:
        result = do_join_query(
            primary_table=data.primary_table,
            max_depth=data.max_depth,
            include_tables=data.include_tables,
            exclude_tables=data.exclude_tables,
            filters=filters,
            order_by=data.order_by,
            page=data.page,
            page_size=data.page_size,
        )
    except ValueError as e:
        raise HttpError(400, str(e))
    
    # 计算执行时间
    execution_time = (time.time() - start_time) * 1000
    
    # 记录查询日志
    user_id = str(request.auth.id) if hasattr(request, 'auth') and request.auth else 'anonymous'
    TableQueryLog.objects.create(
        user_id=user_id,
        table_name=result["join_info"]["primary_table"],
        operation="join_query",
        filters={
            "conditions": filters or [],
            "order_by": data.order_by,
            "max_depth": data.max_depth,
            "joined_tables": result["join_info"]["joined_tables"],
        },
        record_count=len(result["items"]),
        execution_time=execution_time,
    )
    
    logger.info(
        f"联合查询执行: {result['join_info']['primary_table']}, "
        f"关联 {len(result['join_info']['joined_tables'])} 表, "
        f"返回 {len(result['items'])} 条, 耗时 {execution_time:.2f}ms"
    )
    
    # 构建响应
    return JoinQueryResult(
        items=result["items"],
        total=result["total"],
        page=result["page"],
        page_size=result["page_size"],
        join_info=JoinInfo(
            primary_table=result["join_info"]["primary_table"],
            joined_tables=result["join_info"]["joined_tables"],
            join_depth=result["join_info"]["join_depth"],
            total_tables=result["join_info"]["total_tables"],
            has_cycle=result["join_info"]["has_cycle"],
            join_details=[
                JoinTableInfo(
                    table_name=d["table_name"],
                    join_depth=d["join_depth"],
                    source_table=d["source_table"],
                    source_columns=d["source_columns"],
                    target_columns=d["target_columns"],
                )
                for d in result["join_info"]["join_details"]
            ],
        ),
        field_info=[
            FieldInfo(
                alias=f["alias"],
                original_table=f["original_table"],
                original_field=f["original_field"],
                field_type=f["field_type"],
                field_comment=f.get("field_comment", ""),
            )
            for f in result["field_info"]
        ],
    )


@router.get("/table-query/join-preview/{table_name}", response=JoinPreviewOut, tags=["联合查询"], summary="预览表的关联关系")
def get_join_preview(request, table_name: str, max_depth: int = 2):
    """
    预览指定表的外键关联关系
    
    帮助用户了解某个表可以关联哪些表，以及关联的字段信息。
    
    路径参数:
    - table_name: 表名（支持模糊匹配）
    
    查询参数:
    - max_depth: 最大关联深度（可选，1-5级，默认2级）
    
    返回:
    - primary_table: 主表名
    - join_tree: 关联关系树
    - total_related_tables: 总关联表数
    - max_depth: 最大关联深度
    - has_cycle: 是否存在循环引用
    - all_fields: 所有字段列表（带前缀）
    
    AI 调用建议: 在执行联合查询前，先调用此接口了解表的关联关系。
    """
    from core.table_query.join_query_utils import get_join_preview as do_preview
    
    # 限制深度
    max_depth = min(max(1, max_depth), 5)
    
    try:
        result = do_preview(table_name=table_name, max_depth=max_depth)
    except ValueError as e:
        raise HttpError(400, str(e))
    
    return JoinPreviewOut(
        primary_table=result["primary_table"],
        join_tree=[
            JoinTableInfo(
                table_name=d["table_name"],
                join_depth=d["join_depth"],
                source_table=d["source_table"],
                source_columns=d["source_columns"],
                target_columns=d["target_columns"],
            )
            for d in result["join_tree"]
        ],
        total_related_tables=result["total_related_tables"],
        max_depth=result["max_depth"],
        has_cycle=result["has_cycle"],
        all_fields=[
            FieldInfo(
                alias=f["alias"],
                original_table=f["original_table"],
                original_field=f["original_field"],
                field_type=f["field_type"],
            )
            for f in result["all_fields"]
        ],
    )


@router.post("/table-query/join-export", tags=["联合查询"], summary="导出联合查询结果")
def export_join_query(request, data: JoinExportParams):
    """
    导出联合查询结果为 Excel 或 CSV 文件
    
    请求体:
    - primary_table: 主表名或配置名（必填，支持模糊匹配）
    - max_depth: 最大关联深度（可选，1-5级，默认2级）
    - include_tables: 指定要包含的关联表（可选）
    - exclude_tables: 指定要排除的关联表（可选）
    - format: 导出格式 excel/csv（可选，默认 excel）
    - filters: 过滤条件（可选）
    - max_rows: 最大导出行数（可选，默认10000，最大100000）
    
    返回:
    - 文件下载响应
    
    错误:
    - 400: 数据量超过导出上限时返回错误
    """
    from core.table_query.join_query_utils import (
        JoinRelationParser,
        JoinSQLBuilder,
        MAX_PAGE_SIZE,
    )
    
    start_time = time.time()
    
    # 转换过滤条件格式
    filters = None
    if data.filters:
        filters = [{"field": f.field, "operator": f.operator, "value": f.value} for f in data.filters]
    
    try:
        # 解析关联关系
        parser = JoinRelationParser(
            max_depth=data.max_depth,
            include_tables=data.include_tables,
            exclude_tables=data.exclude_tables,
        )
        parse_result = parser.parse(data.primary_table)
        
        # 构建 SQL
        builder = JoinSQLBuilder(
            primary_table=parse_result.primary_table,
            relations=parse_result.relations,
        )
        
        # 构建查询 SQL
        select_clause, field_meta = builder.build_select_fields()
        from_clause = builder.build_from_clause()
        where_clause, where_params = builder.build_where_clause(filters)
        
        # 先查询总数，检查是否超过导出上限
        count_sql_parts = ["SELECT COUNT(*) as cnt", from_clause]
        if where_clause:
            count_sql_parts.append(where_clause)
        count_sql = "\n".join(count_sql_parts)
        
        with connection.cursor() as cursor:
            cursor.execute(count_sql, where_params)
            total_count = cursor.fetchone()[0]
        
        # 检查数据量是否超过导出上限
        if total_count > data.max_rows:
            raise HttpError(
                400, 
                f"数据量过大（{total_count} 条），超过导出上限 {data.max_rows} 条。请添加过滤条件缩小范围后再导出。"
            )
        
        # 构建查询 SQL（使用 max_rows 限制，以防万一）
        sql_parts = [f"SELECT {select_clause}", from_clause]
        if where_clause:
            sql_parts.append(where_clause)
        sql_parts.append(f"LIMIT {data.max_rows}")
        
        sql = "\n".join(sql_parts)
        
        # 执行查询
        with connection.cursor() as cursor:
            cursor.execute(sql, where_params)
            columns = [col[0] for col in cursor.description]
            rows = cursor.fetchall()
        
    except ValueError as e:
        raise HttpError(400, str(e))
    
    # 计算执行时间
    execution_time = (time.time() - start_time) * 1000
    
    # 记录导出日志
    user_id = str(request.auth.id) if hasattr(request, 'auth') and request.auth else 'anonymous'
    TableQueryLog.objects.create(
        user_id=user_id,
        table_name=parse_result.primary_table,
        operation="join_export",
        filters={
            "conditions": filters or [],
            "format": data.format,
            "max_depth": data.max_depth,
            "joined_tables": parse_result.joined_tables,
        },
        record_count=len(rows),
        execution_time=execution_time,
    )
    
    logger.info(
        f"联合查询导出: {parse_result.primary_table}, "
        f"关联 {len(parse_result.relations)} 表, "
        f"导出 {len(rows)} 条, 格式 {data.format}"
    )
    
    # 生成文件名
    filename = f"{parse_result.primary_table}_联合查询"
    
    # 生成导出文件
    if data.format == "csv":
        return _export_csv(filename, columns, rows)
    else:
        return _export_excel(filename, columns, rows)

