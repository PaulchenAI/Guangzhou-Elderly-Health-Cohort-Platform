#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Survey API - 问卷数据管理接口

提供：
- Schema 配置查询
- 问卷数据查询
- 数据导出
- 导入日志查询
"""
import io
import logging
import time
from typing import List, Optional

from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from ninja import Router, Query
from ninja.errors import HttpError
from ninja.pagination import paginate

from common.fu_crud import retrieve
from common.fu_pagination import MyPagination
from core.survey.survey_model import (
    SurveyImportLog,
    SurveyRecord,
    SurveySchemaConfig,
)
from core.survey.survey_schema import (
    SurveyExportIn,
    SurveyImportLogFilters,
    SurveyImportLogSchemaOut,
    SurveyQueryIn,
    SurveyQueryResult,
    SurveyRecordFilters,
    SurveyRecordSchemaOut,
    SurveySchemaConfigFilters,
    SurveySchemaConfigSchemaOut,
    SurveySyncIn,
    SurveySyncOut,
)

logger = logging.getLogger(__name__)
router = Router()


# =============================================================================
# 辅助函数
# =============================================================================

# 子问卷名称映射
SUB_QUESTIONNAIRE_NAMES = {
    "personalityScores": "性格特征评分",
    "personality_scores": "性格特征评分",
    "mentalStateScores": "心理状态评分",
    "mental_state_scores": "心理状态评分",
    "anxietyScore": "焦虑评分",
    "anxiety_score": "焦虑评分",
    "depressionScore": "抑郁评分",
    "depression_score": "抑郁评分",
    "completeness": "完成度统计",
    "scores": "评分详情",
    "statistics": "统计数据",
    "summary": "汇总信息",
    "activities": "活动记录",
    "records": "记录列表",
    "items": "条目列表",
}


def _is_sub_questionnaire(key: str, value) -> bool:
    """
    判断是否为子问卷数据
    
    子问卷数据的特征：
    1. 值是字典类型，且字段名包含特定关键字（scores, Score, statistics 等）
    2. 值是数组类型，且包含多个对象（如 activities）
    """
    # 数组类型：活动列表等
    if isinstance(value, list) and len(value) > 0 and isinstance(value[0], dict):
        # 特定的数组字段名
        array_keywords = ['activities', 'records', 'items', 'entries']
        if key.lower() in array_keywords:
            return True
    
    if not isinstance(value, dict):
        return False
    
    # 根据字段名判断
    key_lower = key.lower()
    sub_keywords = ['scores', 'score', 'statistics', 'summary', 'completeness']
    if any(kw in key_lower for kw in sub_keywords):
        return True
    
    # 根据内容判断（包含评分相关字段）
    score_fields = {'totalScore', 'total_score', 'scores', 'level', 'average', 'count', 'percentage'}
    if score_fields.intersection(value.keys()):
        return True
    
    return False


def _get_sub_questionnaire_name(key: str) -> str:
    """获取子问卷的显示名称"""
    # 先查找映射
    if key in SUB_QUESTIONNAIRE_NAMES:
        return SUB_QUESTIONNAIRE_NAMES[key]
    
    # 将 camelCase 转换为可读格式
    import re
    # 分割驼峰命名
    words = re.sub('([A-Z])', r' \1', key).strip().split()
    return ' '.join(words).title()


# 活动类别映射（数值 -> 文案）
ACTIVITY_CATEGORY_MAP = {
    1: '快步走',
    2: '慢跑',
    3: '骑行',
    4: '健身操/舞蹈',
    5: '球类运动',
    6: '休闲步行',
    7: '园艺活动',
    8: '其他',
    '1': '快步走',
    '2': '慢跑',
    '3': '骑行',
    '4': '健身操/舞蹈',
    '5': '球类运动',
    '6': '休闲步行',
    '7': '园艺活动',
    '8': '其他',
}

# 性别映射
GENDER_MAP = {
    'male': '男',
    'female': '女',
    'Male': '男',
    'Female': '女',
    '男': '男',
    '女': '女',
    1: '男',
    2: '女',
    '1': '男',
    '2': '女',
    0: '女',
    '0': '女',
}


def _build_options_map(fields: list) -> dict:
    """
    从字段定义中构建选项值到标签的映射表
    
    返回格式: {field_name: {value: label, ...}, ...}
    """
    options_map = {}
    for field in fields:
        field_name = field.get("name") or field.get("fieldName")
        options = field.get("options")
        if field_name and options:
            options_map[field_name] = {}
            for opt in options:
                value = opt.get("value")
                label = opt.get("label")
                if value is not None and label:
                    options_map[field_name][value] = label
                    # 同时添加字符串和数字类型的映射
                    if isinstance(value, int):
                        options_map[field_name][str(value)] = label
                    elif isinstance(value, str) and value.isdigit():
                        options_map[field_name][int(value)] = label
    return options_map


def _extract_summary_fields(record) -> dict:
    """
    从问卷记录中提取汇总字段（与 query_records 保持一致）
    """
    summary_fields = {}
    
    # 户外活动：从 summary 提取汇总字段
    if record.survey_type == 'outdoor_activity':
        summary_data = record.survey_data.get('summary', {})
        if summary_data:
            summary_fields = {
                'totalActivities': summary_data.get('totalActivities'),
                'totalDuration': summary_data.get('totalDuration'),
                'averageDuration': summary_data.get('averageDuration'),
                'activityFrequency': summary_data.get('activityFrequency'),
                'mainActivityType': summary_data.get('mainActivityType'),
                'mainVenue': summary_data.get('mainVenue'),
            }
    
    # 空闲活动：从 participationLevel 提取评估字段
    elif record.survey_type == 'leisure_activity':
        participation = record.survey_data.get('participationLevel', {})
        if participation:
            summary_fields = {
                'level': participation.get('level'),
                'diagnosis': participation.get('description'),
            }
    
    # GPM/NRS 疼痛量表：用 painLevel 作为诊断结果
    elif record.survey_type in ['gpm', 'nrs']:
        pain_level = record.survey_data.get('painLevel')
        if pain_level:
            summary_fields['diagnosis'] = pain_level
    
    # 性格特征记录表：从嵌套评分中提取关键字段
    elif record.survey_type == 'personality':
        ms = record.survey_data.get('mentalStateScores', {})
        ps = record.survey_data.get('personalityScores', {})
        
        # 性格特征总分
        if ps:
            summary_fields['personality_total'] = ps.get('totalScore')
        
        # 心理状态总分
        if ms:
            summary_fields['mental_total'] = ms.get('totalScore')
            
            # 焦虑评分
            anxiety = ms.get('anxietyScore', {})
            if anxiety:
                summary_fields['anxiety_score'] = anxiety.get('total')
                anxiety_level = anxiety.get('level', {})
                if isinstance(anxiety_level, dict):
                    summary_fields['anxiety_level'] = anxiety_level.get('level')
            
            # 抑郁评分
            depression = ms.get('depressionScore', {})
            if depression:
                summary_fields['depression_score'] = depression.get('total')
                depression_level = depression.get('level', {})
                if isinstance(depression_level, dict):
                    summary_fields['depression_level'] = depression_level.get('level')
    
    # 通用处理：检查常见的嵌套评估字段
    else:
        for nested_key in ['participationLevel', 'assessmentResult', 'result']:
            nested_data = record.survey_data.get(nested_key, {})
            if isinstance(nested_data, dict) and nested_data:
                if 'level' in nested_data and 'level' not in summary_fields:
                    summary_fields['level'] = nested_data.get('level')
                if 'description' in nested_data and 'diagnosis' not in summary_fields:
                    summary_fields['diagnosis'] = nested_data.get('description')
    
    return summary_fields


def _format_export_value(value, field_name: str, value_mode: str, options_map: dict, field_option: dict = None) -> str:
    """
    根据导出模式格式化值
    
    Args:
        value: 原始值
        field_name: 字段名
        value_mode: 导出模式 ('label' 或 'value')
        options_map: 全局选项映射表
        field_option: 当前字段的选项映射
    
    Returns:
        格式化后的字符串
    """
    if value is None or value == '':
        return ''
    
    # 处理复杂类型（dict/list）
    if isinstance(value, dict):
        # 如果是评估等级对象，提取关键信息
        if 'level' in value:
            if value_mode == 'label':
                return str(value.get('level', '')) + (f"（{value.get('description')}）" if value.get('description') else '')
            else:
                return str(value.get('level', ''))
        return str(value)
    
    if isinstance(value, list):
        if value_mode == 'label':
            # 尝试将列表中的值转换为标签
            labels = []
            for v in value:
                label = _get_value_label(v, field_name, options_map, field_option)
                labels.append(label if label else str(v))
            return '、'.join(labels)
        return '、'.join(str(v) for v in value)
    
    # 导出模式处理
    if value_mode == 'label':
        # 特殊处理：性别字段
        if field_name == 'gender':
            if value in GENDER_MAP:
                return GENDER_MAP[value]
        
        # 特殊处理：活动类别
        if field_name in ['category', 'activityCode', 'mostCommonCategory', 'mainActivityType']:
            if value in ACTIVITY_CATEGORY_MAP:
                return ACTIVITY_CATEGORY_MAP[value]
        
        # 特殊处理：布尔值
        if isinstance(value, bool):
            return '是' if value else '否'
        if field_name == 'isWeekend':
            return '是' if value else '否'
        
        # 尝试从选项映射中获取标签
        label = _get_value_label(value, field_name, options_map, field_option)
        if label:
            return label
    
    # 日期时间格式化
    if isinstance(value, str) and len(value) >= 19 and 'T' in value:
        return value.replace('T', ' ')[:19]
    
    return str(value) if value is not None else ''


def _get_value_label(value, field_name: str, options_map: dict, field_option: dict = None) -> str:
    """
    从选项映射中获取值对应的标签
    """
    # 优先使用字段特定的选项映射
    if field_option:
        if value in field_option:
            return field_option[value]
        # 尝试类型转换后匹配
        if isinstance(value, int) and str(value) in field_option:
            return field_option[str(value)]
        if isinstance(value, str) and value.isdigit():
            int_val = int(value)
            if int_val in field_option:
                return field_option[int_val]
    
    # 使用全局选项映射
    if field_name in options_map:
        field_opts = options_map[field_name]
        if value in field_opts:
            return field_opts[value]
        # 尝试类型转换后匹配
        if isinstance(value, int) and str(value) in field_opts:
            return field_opts[str(value)]
        if isinstance(value, str) and value.isdigit():
            int_val = int(value)
            if int_val in field_opts:
                return field_opts[int_val]
    
    return ''


# =============================================================================
# Schema 配置 API
# =============================================================================

@router.get(
    "/survey/schemas", 
    response=List[SurveySchemaConfigSchemaOut], 
    tags=["问卷管理"],
    summary="获取问卷配置列表（分页）",
)
@paginate(MyPagination)
def list_schemas(request, filters: SurveySchemaConfigFilters = Query(...)):
    """
    获取问卷 Schema 配置列表（分页）
    
    **使用场景**: 查询有哪些问卷类型，获取 schema_id 用于后续数据查询
    **后续操作**: 使用返回的 id 调用 /survey/query 查询问卷数据
    
    查询参数:
    - page: 页码（默认 1）
    - page_size: 每页数量（默认 10）
    - survey_type: 问卷类型标识（如 personality, outdoor_activity）
    - survey_name: 问卷名称（模糊查询，如 "性格"）
    - is_active: 是否激活
    
    返回字段说明:
    - id: Schema ID（用于 query/export 接口）
    - survey_type: 问卷类型标识
    - survey_name: 问卷显示名称
    - fields: 字段配置（定义了可查询的字段）
    """
    return retrieve(request, SurveySchemaConfig, filters)


@router.get(
    "/survey/schemas/all", 
    response=List[SurveySchemaConfigSchemaOut], 
    tags=["问卷管理"],
    summary="获取所有问卷配置",
)
def list_all_schemas(request, is_active: Optional[bool] = True):
    """
    获取所有问卷 Schema 配置（不分页）
    
    查询参数:
    - is_active: 是否只返回激活的配置（默认 True）
    """
    queryset = SurveySchemaConfig.objects.filter(is_deleted=False)
    if is_active is not None:
        queryset = queryset.filter(is_active=is_active)
    return list(queryset)


@router.get(
    "/survey/schemas/{schema_id}", 
    response=SurveySchemaConfigSchemaOut, 
    tags=["问卷管理"],
    summary="获取问卷配置详情",
)
def get_schema(request, schema_id: str):
    """
    获取问卷 Schema 配置详情
    
    路径参数:
    - schema_id: Schema 配置 ID
    """
    return get_object_or_404(SurveySchemaConfig, id=schema_id, is_deleted=False)


@router.get(
    "/survey/schemas/by-type/{survey_type}", 
    response=SurveySchemaConfigSchemaOut, 
    tags=["问卷管理"],
    summary="根据类型获取问卷配置",
)
def get_schema_by_type(request, survey_type: str):
    """
    根据问卷类型获取 Schema 配置
    
    路径参数:
    - survey_type: 问卷类型标识
    """
    return get_object_or_404(
        SurveySchemaConfig, 
        survey_type=survey_type, 
        is_deleted=False
    )


# =============================================================================
# 问卷数据 API
# =============================================================================

@router.get(
    "/survey/records", 
    response=List[SurveyRecordSchemaOut], 
    tags=["问卷管理"],
    summary="获取问卷数据列表（分页）",
)
@paginate(MyPagination)
def list_records(request, filters: SurveyRecordFilters = Query(...)):
    """
    获取问卷数据列表（分页）
    
    查询参数:
    - page: 页码
    - page_size: 每页数量
    - survey_type: 问卷类型
    - survey_id: 问卷ID（模糊查询）
    - patient_name: 患者姓名（模糊查询）
    - start_date: 开始日期
    - end_date: 结束日期
    """
    return retrieve(request, SurveyRecord, filters)


@router.get(
    "/survey/records/{record_id}", 
    response=SurveyRecordSchemaOut, 
    tags=["问卷管理"],
    summary="获取问卷数据详情",
)
def get_record(request, record_id: str):
    """
    获取问卷数据详情
    
    路径参数:
    - record_id: 记录 ID
    """
    return get_object_or_404(SurveyRecord, id=record_id, is_deleted=False)


@router.post(
    "/survey/query", 
    response=SurveyQueryResult, 
    tags=["问卷管理"],
    summary="查询问卷数据",
)
def query_records(request, data: SurveyQueryIn):
    """
    动态查询问卷数据（根据 Schema 配置）
    
    **前置条件**: 需要先调用 /survey/schemas 获取 schema_id
    **典型流程**: 获取配置列表 → 选择问卷类型 → 调用本接口查询数据
    
    请求体:
    - schema_id: Schema 配置 ID（必填，来自 /survey/schemas）
    - page: 页码（默认 1）
    - page_size: 每页数量（默认 10）
    - filters: 过滤条件数组，格式 [{"field": "字段名", "operator": "eq/like/gt/gte/lt/lte", "value": "值"}]
    - order_by: 排序字段（如 "-record_time" 表示按记录时间倒序）
    
    返回值:
    - items: 问卷记录列表（包含主问卷字段和子问卷数据）
    - total: 总记录数
    - page/page_size: 分页信息
    
    **使用示例**:
    查询性格特征问卷: {"schema_id": "xxx", "page": 1, "page_size": 10}
    带条件查询: {"schema_id": "xxx", "filters": [{"field": "patient_name", "operator": "like", "value": "张"}]}
    """
    start_time = time.time()
    
    # 获取 Schema 配置
    schema = get_object_or_404(
        SurveySchemaConfig, 
        id=data.schema_id, 
        is_deleted=False
    )
    
    if not schema.is_active:
        raise HttpError(400, "该问卷配置已禁用")
    
    # 构建查询
    queryset = SurveyRecord.objects.filter(
        survey_type=schema.survey_type,
        is_deleted=False,
    )
    
    # 应用过滤条件
    if data.filters:
        for condition in data.filters:
            field = condition.get("field")
            operator = condition.get("operator", "eq")
            value = condition.get("value")
            
            if not field or value is None:
                continue
            
            # 构建查询条件
            if operator == "eq":
                queryset = queryset.filter(**{field: value})
            elif operator == "like":
                queryset = queryset.filter(**{f"{field}__icontains": value})
            elif operator == "gt":
                queryset = queryset.filter(**{f"{field}__gt": value})
            elif operator == "gte":
                queryset = queryset.filter(**{f"{field}__gte": value})
            elif operator == "lt":
                queryset = queryset.filter(**{f"{field}__lt": value})
            elif operator == "lte":
                queryset = queryset.filter(**{f"{field}__lte": value})
    
    # 排序
    if data.order_by:
        queryset = queryset.order_by(data.order_by)
    else:
        queryset = queryset.order_by("-record_time")
    
    # 分页
    total = queryset.count()
    offset = (data.page - 1) * data.page_size
    records = queryset[offset:offset + data.page_size]
    
    # 构建结果
    items = []
    for record in records:
        # 分离主问卷字段和子问卷数据
        main_fields = {}
        sub_questionnaires = []
        
        # 处理 survey_data
        for key, value in record.survey_data.items():
            if _is_sub_questionnaire(key, value):
                # 这是子问卷数据
                if isinstance(value, list):
                    # 数组类型（如 activities），每个元素作为一条记录
                    sub_questionnaires.append({
                        "name": _get_sub_questionnaire_name(key),
                        "key": key,
                        "type": "array",  # 标记为数组类型
                        "data": value,  # 保留原始数组
                    })
                else:
                    # 字典类型
                    sub_questionnaires.append({
                        "name": _get_sub_questionnaire_name(key),
                        "key": key,
                        "type": "object",
                        "data": value,
                    })
            else:
                # 主问卷字段
                main_fields[key] = value
        
        # 处理 scores（评分相关的子问卷）
        for key, value in record.scores.items():
            if isinstance(value, dict):
                # 检查是否已经在 sub_questionnaires 中（检查 snake_case 和 camelCase 变体）
                # 将 snake_case 转换为 camelCase 进行比较
                key_variants = [key]
                if '_' in key:
                    # snake_case -> camelCase
                    camel = ''.join(word.capitalize() if i > 0 else word for i, word in enumerate(key.split('_')))
                    key_variants.append(camel)
                else:
                    # camelCase -> snake_case
                    import re
                    snake = re.sub('([A-Z])', r'_\1', key).lower().lstrip('_')
                    key_variants.append(snake)
                
                existing = next((sq for sq in sub_questionnaires if sq["key"] in key_variants), None)
                if existing:
                    # 合并到已有的子问卷数据
                    existing["data"].update(value)
                else:
                    # 添加为新的子问卷
                    sub_questionnaires.append({
                        "name": _get_sub_questionnaire_name(key),
                        "key": key,
                        "data": value,
                    })
            else:
                # 简单评分值
                main_fields[key] = value
        
        # 特殊处理：从嵌套数据中提取汇总字段到主表
        summary_fields = {}
        
        # 户外活动：从 summary 提取汇总字段
        if record.survey_type == 'outdoor_activity':
            summary_data = record.survey_data.get('summary', {})
            if summary_data:
                summary_fields = {
                    'totalActivities': summary_data.get('totalActivities'),
                    'totalDuration': summary_data.get('totalDuration'),
                    'averageDuration': summary_data.get('averageDuration'),
                    'activityFrequency': summary_data.get('activityFrequency'),
                    'mainActivityType': summary_data.get('mainActivityType'),
                    'mainVenue': summary_data.get('mainVenue'),
                }
        
        # 空闲活动：从 participationLevel 提取评估字段
        elif record.survey_type == 'leisure_activity':
            participation = record.survey_data.get('participationLevel', {})
            if participation:
                summary_fields = {
                    'level': participation.get('level'),
                    'diagnosis': participation.get('description'),
                }
        
        # GPM/NRS 疼痛量表：用 painLevel 作为诊断结果
        elif record.survey_type in ['gpm', 'nrs']:
            pain_level = record.survey_data.get('painLevel')
            if pain_level:
                summary_fields['diagnosis'] = pain_level
        
        # 性格特征记录表：从嵌套评分中提取关键字段
        elif record.survey_type == 'personality':
            ms = record.survey_data.get('mentalStateScores', {})
            ps = record.survey_data.get('personalityScores', {})
            
            # 性格特征总分
            if ps:
                summary_fields['personality_total'] = ps.get('totalScore')
            
            # 心理状态总分
            if ms:
                summary_fields['mental_total'] = ms.get('totalScore')
                
                # 焦虑评分
                anxiety = ms.get('anxietyScore', {})
                if anxiety:
                    summary_fields['anxiety_score'] = anxiety.get('total')
                    anxiety_level = anxiety.get('level', {})
                    if isinstance(anxiety_level, dict):
                        summary_fields['anxiety_level'] = anxiety_level.get('level')
                
                # 抑郁评分
                depression = ms.get('depressionScore', {})
                if depression:
                    summary_fields['depression_score'] = depression.get('total')
                    depression_level = depression.get('level', {})
                    if isinstance(depression_level, dict):
                        summary_fields['depression_level'] = depression_level.get('level')
        
        # 通用处理：检查常见的嵌套评估字段
        else:
            # 检查 participationLevel、assessmentResult 等常见嵌套对象
            for nested_key in ['participationLevel', 'assessmentResult', 'result']:
                nested_data = record.survey_data.get(nested_key, {})
                if isinstance(nested_data, dict) and nested_data:
                    if 'level' in nested_data and 'level' not in summary_fields:
                        summary_fields['level'] = nested_data.get('level')
                    if 'description' in nested_data and 'diagnosis' not in summary_fields:
                        summary_fields['diagnosis'] = nested_data.get('description')
        
        item = {
            "id": str(record.id),
            "survey_id": record.survey_id,
            "patient_name": record.patient_name,
            "record_time": record.record_time.isoformat() if record.record_time else None,
            **record.patient_info,
            **main_fields,
            **summary_fields,  # 添加汇总字段
            # 子问卷数据供展开行使用
            "_sub_questionnaires": sub_questionnaires if sub_questionnaires else None,
            # 记录者信息
            "recorder_name": record.recorder_info.get("name") if record.recorder_info else None,
        }
        items.append(item)
    
    execution_time = (time.time() - start_time) * 1000
    logger.info(
        f"问卷查询: {schema.survey_type}, "
        f"返回 {len(items)} 条, 耗时 {execution_time:.2f}ms"
    )
    
    return SurveyQueryResult(
        items=items,
        total=total,
        page=data.page,
        page_size=data.page_size,
    )


@router.post("/survey/export", tags=["问卷管理"], summary="导出问卷数据")
def export_records(request, data: SurveyExportIn):
    """
    导出问卷数据（Excel/CSV）
    
    请求体:
    - schema_id: Schema 配置 ID
    - format: 导出格式（excel/csv）
    - value_mode: 导出模式（label=文案, value=数值）
    - filters: 过滤条件
    - max_rows: 最大导出行数
    """
    # 获取 Schema 配置
    schema = get_object_or_404(
        SurveySchemaConfig, 
        id=data.schema_id, 
        is_deleted=False
    )
    
    if not schema.is_active:
        raise HttpError(400, "该问卷配置已禁用")
    
    # 构建查询
    queryset = SurveyRecord.objects.filter(
        survey_type=schema.survey_type,
        is_deleted=False,
    ).order_by("-record_time")
    
    # 应用过滤条件
    if data.filters:
        for condition in data.filters:
            field = condition.get("field")
            operator = condition.get("operator", "eq")
            value = condition.get("value")
            
            if not field or value is None:
                continue
            
            if operator == "eq":
                queryset = queryset.filter(**{field: value})
            elif operator == "like":
                queryset = queryset.filter(**{f"{field}__icontains": value})
    
    # 限制行数
    records = queryset[:data.max_rows]
    
    # 获取字段配置
    fields = schema.get_visible_fields()
    
    # 创建选项值到标签的映射表（用于文案导出模式）
    options_map = _build_options_map(fields)
    
    # 获取导出模式
    value_mode = data.value_mode  # 'label' 或 'value'
    
    # 跳过的字段（已经作为固定列或不需要导出）
    skip_fields = {
        'surveyId', 'survey_id', 'patientName', 'patient_name',
        'recordTime', 'record_time', 'createdAt', 'created_at',
        'age', 'gender', 'phone', 'remarks',
        'patientId', 'recorderName', 'recorderPosition',
        'recorderInstitution', 'recorderContact',
    }
    
    # 构建数据行 - 先添加固定字段
    headers = ["问卷ID", "患者姓名", "年龄", "性别", "记录时间"]
    field_names = ["survey_id", "patient_name", "age", "gender", "record_time"]
    field_options = [None, None, None, None, None]  # 对应每个字段的选项映射
    
    # 添加 Schema 定义的字段（使用 name/title 或 fieldName/label）
    for f in fields:
        field_name = f.get("name") or f.get("fieldName")
        field_title = f.get("title") or f.get("label") or field_name
        if field_name and field_name not in field_names and field_name not in skip_fields:
            headers.append(field_title)
            field_names.append(field_name)
            # 记录该字段的选项映射
            field_options.append(options_map.get(field_name))
    
    # 特殊处理：根据问卷类型添加汇总字段
    if schema.survey_type == 'outdoor_activity':
        # 户外活动汇总字段
        summary_fields = [
            ('totalActivities', '活动总数'),
            ('totalDuration', '总时长(分钟)'),
            ('averageDuration', '平均时长(分钟)'),
            ('activityFrequency', '活动频率'),
            ('mainActivityType', '主要活动类型'),
            ('mainVenue', '主要场所'),
        ]
        for field_name, field_title in summary_fields:
            if field_name not in field_names:
                headers.append(field_title)
                field_names.append(field_name)
                field_options.append(None)
    
    elif schema.survey_type == 'personality':
        # 性格特征汇总字段
        personality_fields = [
            ('personality_total', '性格总分'),
            ('mental_total', '心理总分'),
            ('anxiety_score', '焦虑分数'),
            ('anxiety_level', '焦虑等级'),
            ('depression_score', '抑郁分数'),
            ('depression_level', '抑郁等级'),
        ]
        for field_name, field_title in personality_fields:
            if field_name not in field_names:
                headers.append(field_title)
                field_names.append(field_name)
                field_options.append(None)
    
    else:
        # 其他问卷添加诊断/结果列
        result_fields = [
            ('diagnosis', '诊断结果'),
            ('level', '评估等级'),
        ]
        for field_name, field_title in result_fields:
            if field_name not in field_names:
                headers.append(field_title)
                field_names.append(field_name)
                field_options.append(None)
    
    # 添加记录者信息
    headers.append("记录者")
    field_names.append("recorder_name")
    field_options.append(None)
    
    # 检查是否需要导出子表单（户外活动等）
    has_activities = schema.survey_type == 'outdoor_activity'
    
    if has_activities:
        # 添加子活动字段头
        activity_headers = ["活动序号", "活动类别", "活动代码", "开始时间", "时长(分钟)", "场所", "路线", "是否周末", "使用物品"]
        activity_fields = ["activity_index", "category", "activityCode", "startTime", "duration", "venue", "route", "isWeekend", "itemsUsed"]
        headers.extend(activity_headers)
        field_names.extend(activity_fields)
        field_options.extend([None] * len(activity_headers))
    
    rows = []
    for record in records:
        # 提取汇总字段（与前端 query_records 保持一致）
        summary_fields_data = _extract_summary_fields(record)
        
        # 合并所有数据源
        all_data = {
            "survey_id": record.survey_id,
            "patient_name": record.patient_name,
            "record_time": record.record_time.strftime("%Y-%m-%d %H:%M:%S") if record.record_time else "",
            "recorder_name": record.recorder_info.get("name", "") if record.recorder_info else "",
            **record.patient_info,
            **record.survey_data,
            **record.scores,
            **summary_fields_data,
        }
        
        if has_activities:
            # 户外活动：展开 activities 数组为多行
            activities = record.survey_data.get('activities', [])
            if activities:
                for idx, activity in enumerate(activities):
                    row = []
                    for i, field_name in enumerate(field_names):
                        if field_name == "activity_index":
                            row.append(str(idx + 1))
                        elif field_name in ["category", "activityCode", "startTime", "duration", "venue", "route", "isWeekend", "itemsUsed"]:
                            value = activity.get(field_name, "")
                            row.append(_format_export_value(value, field_name, value_mode, options_map, field_options[i] if i < len(field_options) else None))
                        else:
                            value = all_data.get(field_name, "")
                            row.append(_format_export_value(value, field_name, value_mode, options_map, field_options[i] if i < len(field_options) else None))
                    rows.append(row)
            else:
                # 没有活动记录，导出一行空数据
                row = []
                for i, field_name in enumerate(field_names):
                    if field_name in ["activity_index", "category", "activityCode", "startTime", "duration", "venue", "route", "isWeekend", "itemsUsed"]:
                        row.append("")
                    else:
                        value = all_data.get(field_name, "")
                        row.append(_format_export_value(value, field_name, value_mode, options_map, field_options[i] if i < len(field_options) else None))
                rows.append(row)
        else:
            # 普通问卷：一条记录一行
            row = []
            for i, field_name in enumerate(field_names):
                value = all_data.get(field_name, "")
                row.append(_format_export_value(value, field_name, value_mode, options_map, field_options[i] if i < len(field_options) else None))
            rows.append(row)
    
    mode_label = "文案" if value_mode == "label" else "数值"
    logger.info(f"导出问卷数据: {schema.survey_type}, 共 {len(rows)} 条, 模式: {mode_label}")
    
    # 生成文件（文件名包含导出模式）
    filename = f"{schema.survey_name}_{mode_label}"
    if data.format == "csv":
        return _export_csv(filename, headers, rows)
    else:
        return _export_excel(filename, headers, rows)


def _export_csv(filename: str, headers: List[str], rows: List[list]) -> HttpResponse:
    """导出为 CSV 格式"""
    import csv
    
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    
    response = HttpResponse(
        output.getvalue(),
        content_type='text/csv; charset=utf-8-sig'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    return response


def _export_excel(filename: str, headers: List[str], rows: List[list]) -> HttpResponse:
    """导出为 Excel 格式"""
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HttpError(500, "Excel 导出需要安装 openpyxl 库")
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = filename[:31]
    
    # 写入表头
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = openpyxl.styles.Font(bold=True)
    
    # 写入数据
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # 调整列宽
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15
    
    # 保存到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response


# =============================================================================
# 导入日志 API
# =============================================================================

@router.get(
    "/survey/import-logs", 
    response=List[SurveyImportLogSchemaOut], 
    tags=["问卷管理"],
    summary="获取导入日志列表（分页）",
)
@paginate(MyPagination)
def list_import_logs(request, filters: SurveyImportLogFilters = Query(...)):
    """
    获取导入日志列表（分页）
    
    查询参数:
    - page: 页码
    - page_size: 每页数量
    - batch_id: 批次ID（模糊查询）
    - survey_type: 问卷类型
    - status: 状态
    - trigger_type: 触发类型
    """
    return retrieve(request, SurveyImportLog, filters)


@router.get(
    "/survey/import-logs/{log_id}", 
    response=SurveyImportLogSchemaOut, 
    tags=["问卷管理"],
    summary="获取导入日志详情",
)
def get_import_log(request, log_id: str):
    """
    获取导入日志详情
    
    路径参数:
    - log_id: 日志 ID
    """
    return get_object_or_404(SurveyImportLog, id=log_id)


# =============================================================================
# 手动同步 API
# =============================================================================

@router.post(
    "/survey/sync",
    response=SurveySyncOut,
    tags=["问卷管理"],
    summary="触发问卷数据同步",
)
def trigger_sync(request, data: SurveySyncIn):
    """
    手动触发问卷数据同步
    
    请求体:
    - incremental: 是否增量同步（默认 True）
    - survey_type: 问卷类型（为空则同步全部）
    
    返回:
    - batch_id: 批次ID
    - total: 总处理数量
    - success: 成功数量
    - skipped: 跳过数量
    - failed: 失败数量
    - message: 结果消息
    """
    from core.survey.survey_service import SurveyService
    
    try:
        service = SurveyService()
        result = service.import_data(
            survey_type=data.survey_type,
            incremental=data.incremental,
            trigger_type="manual",  # 标记为手动触发
        )
        
        # 构建结果消息
        message = f"同步完成：成功 {result['success']} 条"
        if result['skipped'] > 0:
            message += f"，跳过 {result['skipped']} 条"
        if result['failed'] > 0:
            message += f"，失败 {result['failed']} 条"
        
        return SurveySyncOut(
            batch_id=result['batch_id'],
            total=result['total'],
            success=result['success'],
            skipped=result['skipped'],
            failed=result['failed'],
            message=message,
        )
    except Exception as e:
        logger.error(f"手动同步失败: {e}")
        raise HttpError(500, f"同步失败: {str(e)}")

